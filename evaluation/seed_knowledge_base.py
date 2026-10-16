"""Idempotently index the evaluation policy corpus into the demo tenant."""

from html import escape
from pathlib import Path

from sqlalchemy import or_

from app.db import SessionLocal
from app.models.document import Document
from app.models.document_permission import DocumentPermission
from app.rag.loaders import load_text
from app.models.user import User
from app.rag.ingestion import delete_document_vectors, ingest_document
from evaluation.knowledge_base import KNOWLEDGE_BASE_DOCUMENTS

FILES_DIR = Path(__file__).with_name("knowledge_base_files")
TITLE_ALIASES = {
    "RAGVault Security Architecture": ("ragvault-security-architecture",),
    "RAGVault Engineering Architecture": ("Engineering Architecture",),
    "Employee Records Handling Standard": ("Confidential HR Policy",),
    "RAGVault Product Overview": ("RAGVault Test Document",),
}
LEGACY_DUPLICATES = (("upload_test", "upload_test.docx"),)


def materialize_source(item: dict[str, str]) -> str:
    if "source_path" in item:
        source = (Path(__file__).parent / item["source_path"]).read_text(encoding="utf-8")
        return _append_supplement(source, item)

    file_name = item["file_name"]
    content = _append_supplement(item["content"], item)
    format_name = item.get("format")
    if format_name is None:
        FILES_DIR.mkdir(parents=True, exist_ok=True)
        path = FILES_DIR / file_name
        path.write_text(content, encoding="utf-8")
        return path.read_text(encoding="utf-8")

    FILES_DIR.mkdir(parents=True, exist_ok=True)
    path = FILES_DIR / file_name
    if format_name == "pdf":
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        styles = getSampleStyleSheet()
        story = []
        for line in content.splitlines():
            if line.strip():
                story.append(Paragraph(escape(line), styles["BodyText"]))
            else:
                story.append(Spacer(1, 6))
        SimpleDocTemplate(str(path), pagesize=letter).build(story)
    elif format_name == "docx":
        from docx import Document as WordDocument

        document = WordDocument()
        for line in content.splitlines():
            document.add_paragraph(line)
        document.save(path)
    else:
        raise ValueError(f"Unsupported generated knowledge-base format: {format_name}")

    return load_text(str(path), path.suffix)


def _append_supplement(content: str, item: dict[str, str]) -> str:
    supplement = item.get("supplement", "").strip()
    return f"{content.rstrip()}\n\n{supplement}\n" if supplement else content


def _remove_verified_legacy_duplicates(db, tenant_id: str) -> int:
    removed = 0
    for title, file_name in LEGACY_DUPLICATES:
        matches = (
            db.query(Document)
            .filter(
                Document.tenant_id == tenant_id,
                Document.title == title,
                Document.file_name == file_name,
            )
            .order_by(Document.id)
            .all()
        )
        if len(matches) < 2:
            continue

        primary = matches[0]
        for duplicate in matches[1:]:
            if duplicate.content != primary.content:
                raise RuntimeError(
                    f"Refusing to consolidate non-identical duplicate candidates "
                    f"{primary.id} and {duplicate.id}"
                )
            permission_count = (
                db.query(DocumentPermission)
                .filter(DocumentPermission.document_id == duplicate.id)
                .count()
            )
            if permission_count:
                raise RuntimeError(
                    f"Refusing to consolidate duplicate document {duplicate.id} "
                    "with existing permissions"
                )
            delete_document_vectors(duplicate.id)
            db.delete(duplicate)
            removed += 1
    if removed:
        db.commit()
    return removed


def seed_knowledge_base() -> dict[str, int]:
    db = SessionLocal()
    inserted = 0
    updated = 0
    skipped = 0
    try:
        admin = (
            db.query(User)
            .filter(User.email == "admin@ragvault.local")
            .first()
        )
        if admin is None:
            raise RuntimeError("Evaluation admin user is unavailable")

        removed_duplicates = _remove_verified_legacy_duplicates(
            db,
            admin.tenant_id,
        )

        for item in KNOWLEDGE_BASE_DOCUMENTS:
            candidate_titles = (
                item["title"],
                *TITLE_ALIASES.get(item["title"], ()),
            )
            candidate_file_names = item.get("legacy_file_names", ())
            match_conditions = [
                Document.title.in_(candidate_titles),
            ]
            if candidate_file_names:
                match_conditions.append(
                    Document.file_name.in_(candidate_file_names)
                )
            matches = (
                db.query(Document)
                .filter(
                    Document.tenant_id == admin.tenant_id,
                    or_(*match_conditions),
                )
                .order_by(Document.id)
                .all()
            )
            if not matches:
                extracted_content = materialize_source(item)
                if not extracted_content.strip():
                    raise ValueError(
                        f"Generated source document is empty: {item['file_name']}"
                    )
                document = Document(
                    title=item["title"],
                    file_name=item["file_name"],
                    content=extracted_content,
                    tenant_id=admin.tenant_id,
                    classification="Internal",
                )
                db.add(document)
                db.commit()
                db.refresh(document)
                ingest_document(db=db, document=document)
                inserted += 1
                continue

            primary = next(
                (
                    document
                    for alias in TITLE_ALIASES.get(item["title"], ())
                    for document in matches
                    if document.title == alias
                ),
                None,
            )
            if primary is None and candidate_file_names:
                primary = next(
                    (
                        document
                        for document in matches
                        if document.file_name in candidate_file_names
                    ),
                    None,
                )
            if primary is None:
                primary = matches[0]
            duplicates = [document for document in matches if document.id != primary.id]
            for duplicate in duplicates:
                permission_count = (
                    db.query(DocumentPermission)
                    .filter(DocumentPermission.document_id == duplicate.id)
                    .count()
                )
                if permission_count:
                    raise RuntimeError(
                        f"Refusing to consolidate document {duplicate.id} with "
                        "existing permissions"
                    )
                delete_document_vectors(duplicate.id)
                db.delete(duplicate)

            new_content = materialize_source(item)
            if not new_content.strip():
                raise ValueError(f"Source document is empty: {item['file_name']}")
            if (
                primary.title == item["title"]
                and primary.file_name == item["file_name"]
                and primary.content == new_content
                and not duplicates
            ):
                skipped += 1
                continue

            delete_document_vectors(primary.id)
            primary.title = item["title"]
            primary.file_name = item["file_name"]
            primary.content = new_content
            primary.classification = "Internal"
            db.commit()
            ingest_document(db=db, document=primary)
            updated += 1

        documents = (
            db.query(Document)
            .filter(Document.tenant_id == admin.tenant_id)
            .order_by(Document.id)
            .all()
        )
        titles = [document.title for document in documents]
        if len(documents) != 35 or len(set(titles)) != 35:
            raise RuntimeError(
                f"Expected exactly 35 unique evaluation documents, found "
                f"{len(documents)} records and {len(set(titles))} unique titles"
            )
        short_documents = [
            document.title
            for document in documents
            if sum(bool(line.strip()) for line in (document.content or "").splitlines())
            <= 25
        ]
        if short_documents:
            raise RuntimeError(
                "Evaluation documents must each contain more than 25 meaningful "
                f"lines: {', '.join(short_documents)}"
            )
        return {
            "inserted": inserted,
            "updated": updated,
            "skipped": skipped,
            "removed_duplicates": removed_duplicates,
            "document_count": len(documents),
        }
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    print(seed_knowledge_base())
